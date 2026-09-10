#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
JSON to Excel Importer
将 Assets/Configs 下的 JSON 写回 Attribute.xlsm 的同名 sheet（export_to_json.py 的反向）。

表结构约定（与 export_to_json.py 一致）：
    第 1 行 = 列名（JSON 字段名）
    第 2 行 = 类型（string / int / float / enum）
    第 3 行起 = 数据

用法:
    python import_from_json.py                       # 只导入 AttributeValueConfig
    python import_from_json.py AttributeValueConfig CostOffsetConfig ...
"""

import openpyxl
import json
import sys
from pathlib import Path

ROW_VAR, ROW_TYPE, ROW_DATA_START = 1, 2, 3

SCRIPT_DIR = Path(__file__).parent.resolve()
XLSM = SCRIPT_DIR / "Attribute.xlsm"
JSON_DIR = SCRIPT_DIR.parent / "Assets" / "Configs"


def import_sheet(wb, sheet_name: str, records: list) -> None:
    if sheet_name not in wb.sheetnames:
        raise KeyError(f"xlsm 中不存在 sheet '{sheet_name}'")
    ws = wb[sheet_name]

    # 列映射：第 1 行列名 -> 列号
    col_map = {}
    for col in range(1, ws.max_column + 1):
        name = ws.cell(row=ROW_VAR, column=col).value
        if name and str(name).strip():
            col_map[str(name).strip()] = col

    missing = [k for k in records[0].keys() if k not in col_map] if records else []
    if missing:
        raise KeyError(f"sheet 缺少列: {missing}（现有列: {list(col_map)}）")

    # 类型标记
    col_type = {}
    for name, col in col_map.items():
        col_type[name] = str(ws.cell(row=ROW_TYPE, column=col).value or "string").strip().lower()

    # 清空旧数据（含残留样式之外的一切旧值）
    if ws.max_row >= ROW_DATA_START:
        ws.delete_rows(ROW_DATA_START, ws.max_row - ROW_DATA_START + 1)

    for i, rec in enumerate(records):
        row = ROW_DATA_START + i
        for name, value in rec.items():
            col = col_map[name]
            if value is None:
                continue
            t = col_type[name]
            if t == "int":
                try:
                    value = int(value)
                except (ValueError, TypeError):
                    pass
            elif t == "float":
                try:
                    value = float(value)
                except (ValueError, TypeError):
                    pass
            ws.cell(row=row, column=col, value=value)

    print(f"Sheet '{sheet_name}': 写入 {len(records)} 行")


def main():
    names = sys.argv[1:] or ["AttributeValueConfig"]

    wb = openpyxl.load_workbook(XLSM, keep_vba=True)
    for name in names:
        json_file = JSON_DIR / f"{name}.json"
        if not json_file.exists():
            raise FileNotFoundError(json_file)
        with open(json_file, encoding="utf-8") as f:
            records = json.load(f)
        import_sheet(wb, name, records)

    wb.save(XLSM)
    print(f"已保存: {XLSM}")


if __name__ == "__main__":
    main()
